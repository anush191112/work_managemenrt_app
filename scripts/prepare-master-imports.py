#!/usr/bin/env python3
"""Create Supabase Table Editor CSV files from the approved master workbooks."""

from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import uuid
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ARCHITECT_COLUMNS = ["state", "serial_no", "name", "phone", "email", "address", "birthdate"]
BIRTHDAY_COLUMNS = ["name", "phone", "email", "address", "birthdate"]
CAMPAIGN_COLUMNS = [
    "name", "campaign_type", "source", "account", "state", "serial_start", "serial_end",
    "campaign_date", "created_by", "assigned_to", "is_shared", "campaign_completed_on",
    "total_attempted", "total_delivered_first", "retarget_name", "retarget_due_date",
    "retarget_completed_on", "retarget_attempted", "retarget_delivered", "total_delivered_final",
    "percentage", "total_failed", "total_replies", "total_instagram_clicks",
    "total_whatsapp_channel_clicks", "total_catalog_requests", "total_inquiries", "total_leads", "stage",
]
LEAD_COLUMNS = ["id", "name", "phone", "account", "purpose", "created_by", "touch_count", "engagement_tag", "last_marked_by"]
INTERACTION_COLUMNS = ["lead_id", "reported_name", "account", "purpose", "marked_by", "source"]


def value(row: dict[str, object], header: str) -> object:
    return row.get(header, "")


def clean_text(raw: object) -> str:
    return str(raw or "").strip()


def nullable_text(raw: object) -> str:
    return clean_text(raw)


def normalize_phone(raw: object) -> str:
    source = clean_text(raw)
    if re.fullmatch(r"\d+\.0+", source):
        source = source.split(".")[0]
    digits = re.sub(r"\D", "", source)
    if len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
    if len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]
    if not re.fullmatch(r"[6-9]\d{9}", digits) or len(set(digits)) == 1:
        return ""
    return digits


def normalize_email(raw: object) -> str:
    email = clean_text(raw).lower()
    return email if re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email) else ""


def number(raw: object) -> str:
    if raw is None or raw == "":
        return ""
    try:
        parsed = float(raw)
    except (TypeError, ValueError):
        return ""
    return str(int(parsed)) if parsed.is_integer() else str(parsed)


def non_negative_number(raw: object) -> str:
    value = number(raw)
    return value if not value or float(value) >= 0 else ""


def as_date(raw: object) -> str:
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    if isinstance(raw, date):
        return raw.isoformat()
    return ""


def birthday_date(raw: object) -> str:
    if isinstance(raw, (datetime, date)):
        return f"2000-{raw.month:02d}-{raw.day:02d}"
    match = re.fullmatch(r"\s*(\d{1,2})[-/]?(\d{1,2})(?:[-/]\d{2,4})?\s*", clean_text(raw))
    if not match:
        return ""
    day, month = map(int, match.groups())
    try:
        return date(2000, month, day).isoformat()
    except ValueError:
        return ""


def records(path: Path, sheet: str | None = None) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet] if sheet else workbook.active
    row_iter = worksheet.iter_rows(values_only=True)
    raw_headers = [clean_text(cell) for cell in next(row_iter)]
    seen_headers: Counter[str] = Counter()
    headers = []
    for header in raw_headers:
        seen_headers[header] += 1
        headers.append(header if seen_headers[header] == 1 else f"{header}_{seen_headers[header] - 1}")
    output = []
    for line, cells in enumerate(row_iter, start=2):
        if not any(cell not in (None, "") for cell in cells):
            continue
        output.append({"_line": line, **{headers[index]: cells[index] if index < len(cells) else "" for index in range(len(headers))}})
    return output


def write_csv(path: Path, columns: list[str], rows: list[dict[str, object]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def split_csv(directory: Path, prefix: str, columns: list[str], rows: list[dict[str, object]], size: int = 10_000) -> list[str]:
    names = []
    for index in range(0, len(rows), size):
        name = f"{prefix}_{index // size + 1:02d}.csv"
        write_csv(directory / name, columns, rows[index:index + size])
        names.append(name)
    return names


def state_from_campaign(name: str) -> str:
    lowered = name.lower()
    state_markers = {
        "kerala": "Kerala", "kr": "Kerala", "kl": "Kerala",
        "gujarat": "Gujarat", "gujrat": "Gujarat", "gj": "Gujarat",
        "maharashtra": "Maharashtra", "mh": "Maharashtra", "mum": "Maharashtra", "nvmun": "Maharashtra",
        "telangana": "Telangana", "tg": "Telangana", "delhi": "Delhi",
        "karnataka": "Karnataka", "ktka": "Karnataka", "ka": "Karnataka", "tn": "Tamil Nadu",
        "rajasthan": "Rajasthan", "raj": "Rajasthan", "andhra": "Andhra Pradesh", "andp": "Andhra Pradesh",
        "madhya": "Madhya Pradesh", "mp": "Madhya Pradesh", "chhattisgarh": "Chhattisgarh", "chg": "Chhattisgarh",
        "haryana": "Haryana", "hr": "Haryana", "west bengal": "West Bengal", "wb": "West Bengal",
        "bihar": "Bihar", "uttar": "Uttar Pradesh", "up": "Uttar Pradesh", "odisha": "Odisha",
        "puducherry": "Puducherry", "goa": "Goa", "jharkhand": "Jharkhand",
        "chndgrh": "Chandigarh", "pb": "Punjab", "uk": "Uttarakhand", "hp": "Himachal Pradesh",
        "d&d": "Dadra and Nagar Haveli and Daman and Diu", "assam": "Assam", "j&k": "Jammu and Kashmir",
        "mhsub": "Maharashtra", "tl": "Telangana",
    }
    for marker, state in state_markers.items():
        if re.search(rf"(?:^|[^a-z]){re.escape(marker)}(?:[^a-z]|$)", lowered):
            return state
    return ""


def range_from_campaign(name: str) -> tuple[str, str]:
    match = re.search(r"(?<!\d)(\d{1,6})\s*-\s*(\d{1,6})(?!\d)", name)
    return match.groups() if match else ("", "")


def account(raw: object, campaign_type: str) -> str:
    source = clean_text(raw).lower().replace(".0", "")
    mapping = {
        "9928781655": "whatsapp_9928781655",
        "9928501655": "whatsapp_9928501655",
        "stonekingexports@gmail.com": "email_stonekingexports",
        "stonekingexportindia@gmail.com": "email_stonekingexportindia",
    }
    result = mapping.get(source, "")
    return result if result.startswith(campaign_type) else ""


def tag(touch_count: int) -> str:
    if touch_count >= 4:
        return "fan"
    if touch_count == 3:
        return "regular"
    if touch_count == 2:
        return "interested"
    return "curious"


def prepare_architects(source: Path, destination: Path, rejects: list[dict[str, object]]) -> list[str]:
    output: list[dict[str, object]] = []
    state_serials: set[tuple[str, str]] = set()
    name_phones: set[tuple[str, str]] = set()
    for row in records(source, "Cleaned Data"):
        state, serial_no, name = clean_text(value(row, "State")), number(value(row, "Serial Number")), clean_text(value(row, "Full Name"))
        phone, email = normalize_phone(value(row, "Mobile Number")), normalize_email(value(row, "Email Address"))
        if not state or not serial_no or not name:
            rejects.append({"source": "architects", "line": row["_line"], "reason": "Missing state, serial number, or full name"})
            continue
        if (state, serial_no) in state_serials:
            rejects.append({"source": "architects", "line": row["_line"], "reason": "Duplicate state and serial number"})
            continue
        if phone and (name, phone) in name_phones:
            rejects.append({"source": "architects", "line": row["_line"], "reason": "Duplicate name and phone pair"})
            continue
        if not phone and not email:
            rejects.append({"source": "architects", "line": row["_line"], "reason": "No valid phone or email"})
            continue
        state_serials.add((state, serial_no))
        if phone:
            name_phones.add((name, phone))
        output.append({"state": state, "serial_no": serial_no, "name": name, "phone": phone, "email": email, "address": nullable_text(value(row, "Address")), "birthdate": birthday_date(value(row, "Birthday"))})
    return split_csv(destination, "architects", ARCHITECT_COLUMNS, output)


def prepare_birthdays(source: Path, destination: Path, rejects: list[dict[str, object]]) -> str:
    output, seen = [], set()
    for row in records(source):
        name, phone, birthdate = clean_text(value(row, "Name")), normalize_phone(value(row, "Phone Number")), birthday_date(value(row, "Birthday"))
        if not name or not phone or not birthdate:
            rejects.append({"source": "birthday_contacts", "line": row["_line"], "reason": "Name, valid phone, and birthday are required"})
            continue
        if phone in seen:
            rejects.append({"source": "birthday_contacts", "line": row["_line"], "reason": "Duplicate phone"})
            continue
        seen.add(phone)
        output.append({"name": name, "phone": phone, "email": "", "address": "", "birthdate": birthdate})
    name = "birthday_contacts.csv"
    write_csv(destination / name, BIRTHDAY_COLUMNS, output)
    return name


def campaign_row(name: str, campaign_type: str, account_value: str, campaign_date: str, owner_id: str, data: dict[str, object]) -> tuple[dict[str, object] | None, str]:
    state = state_from_campaign(name)
    serial_start, serial_end = range_from_campaign(name)
    source = "birthday" if "bday" in name.lower() or "birthday" in name.lower() else "state_wise"
    if source == "state_wise" and (not state or not serial_start or not serial_end):
        return None, "Could not safely derive state and serial range from campaign name"
    if not account_value or not campaign_date:
        return None, "Missing supported account or campaign date"
    result = {"name": name, "campaign_type": campaign_type, "source": source, "account": account_value, "state": state if source == "state_wise" else "", "serial_start": serial_start if source == "state_wise" else "", "serial_end": serial_end if source == "state_wise" else "", "campaign_date": campaign_date, "created_by": owner_id, "assigned_to": "", "is_shared": "true", "campaign_completed_on": "", "total_attempted": "", "total_delivered_first": "", "retarget_name": "", "retarget_due_date": "", "retarget_completed_on": "", "retarget_attempted": "", "retarget_delivered": "", "total_delivered_final": "", "percentage": "", "total_failed": "", "total_replies": "", "total_instagram_clicks": "", "total_whatsapp_channel_clicks": "", "total_catalog_requests": "", "total_inquiries": "", "total_leads": "", "stage": "complete"}
    result.update(data)
    return result, ""


def prepare_campaigns(whatsapp_source: Path, email_source: Path, destination: Path, owner_id: str, rejects: list[dict[str, object]]) -> list[str]:
    whatsapp, email = [], []
    for row in records(whatsapp_source):
        name = clean_text(value(row, "Name of Campaign"))
        item, reason = campaign_row(name, "whatsapp", account(value(row, "Account Used"), "whatsapp"), as_date(value(row, "Date of Campaign")), owner_id, {
            "campaign_completed_on": as_date(value(row, "Completed On")), "total_attempted": non_negative_number(value(row, "Total Attempted")), "total_delivered_first": non_negative_number(value(row, "Total Delivered(FIrst)")), "retarget_name": nullable_text(value(row, "Retarget Campaign Name")), "retarget_due_date": as_date(value(row, "Date of Retarget Campaign")), "retarget_completed_on": as_date(value(row, "Completed On_1")), "retarget_attempted": non_negative_number(value(row, "Total Retarget Attempted")), "retarget_delivered": non_negative_number(value(row, "Total Delivered(Retarget)")), "total_delivered_final": non_negative_number(value(row, "Total Delivered(Final)")), "percentage": number(value(row, "PERCENTAGE")), "total_failed": non_negative_number(value(row, "Total Failed((First +Retarget)")), "total_replies": non_negative_number(value(row, "Total Replies")), "total_instagram_clicks": non_negative_number(value(row, "Total Instagram Clicks")), "total_whatsapp_channel_clicks": non_negative_number(value(row, "Total Whatsapp Channel Clicks")), "total_catalog_requests": non_negative_number(value(row, "Total Catalog Requests")), "total_inquiries": non_negative_number(value(row, "Total Inquiries")), "total_leads": non_negative_number(value(row, "Total Leads")),
        })
        if item: whatsapp.append(item)
        else: rejects.append({"source": "whatsapp_campaigns", "line": row["_line"], "reason": reason, "name": name})
    for row in records(email_source):
        name = clean_text(value(row, "Name of Campaign"))
        item, reason = campaign_row(name, "email", account(value(row, "Account Used"), "email"), as_date(value(row, "Date of Campaign")), owner_id, {
            "total_attempted": non_negative_number(value(row, "Total Attempted")), "total_delivered_final": non_negative_number(value(row, "Total Delivered")), "total_failed": non_negative_number(value(row, "Total Bounces")), "total_leads": non_negative_number(value(row, "Total Leads")),
        })
        if item: email.append(item)
        else: rejects.append({"source": "email_campaigns", "line": row["_line"], "reason": reason, "name": name})
    write_csv(destination / "campaigns_whatsapp.csv", CAMPAIGN_COLUMNS, whatsapp)
    write_csv(destination / "campaigns_email.csv", CAMPAIGN_COLUMNS, email)
    return ["campaigns_whatsapp.csv", "campaigns_email.csv"]


def prepare_leads(source: Path, destination: Path, owner_id: str, rejects: list[dict[str, object]]) -> list[str]:
    by_phone: dict[str, list[dict[str, object]]] = defaultdict(list)
    purpose_map = {"requested for catalog": "requested_for_catalog", "instagram/ whatsapp channel": "instagram_whatsapp_channel", "inquiry": "inquiry", "casual reply": "casual_reply"}
    for row in records(source):
        name, phone = clean_text(value(row, "Name")), normalize_phone(value(row, "Mobile Number"))
        purpose = purpose_map.get(clean_text(value(row, "Interaction")).lower(), "")
        account_value = account(value(row, "Account Used"), "whatsapp")
        if not name or not phone or not purpose or not account_value:
            rejects.append({"source": "historical_leads", "line": row["_line"], "reason": "Name, valid phone, supported account, and known interaction are required"})
            continue
        by_phone[phone].append({"name": name, "account": account_value, "purpose": purpose})
    leads, interactions = [], []
    for phone, events in sorted(by_phone.items()):
        lead_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"work-management-lead:{phone}"))
        first = events[0]
        leads.append({"id": lead_id, "name": first["name"], "phone": phone, "account": first["account"], "purpose": first["purpose"], "created_by": owner_id, "touch_count": str(len(events)), "engagement_tag": tag(len(events)), "last_marked_by": owner_id})
        interactions.extend({"lead_id": lead_id, "reported_name": event["name"], "account": event["account"], "purpose": event["purpose"], "marked_by": owner_id, "source": "historical_import"} for event in events)
    write_csv(destination / "leads.csv", LEAD_COLUMNS, leads)
    write_csv(destination / "lead_interactions.csv", INTERACTION_COLUMNS, interactions)
    return ["leads.csv", "lead_interactions.csv"]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--owner-id", required=True, help="Existing admin user UUID used as the historical-import owner.")
    parser.add_argument("--output", default="prepared-imports")
    parser.add_argument("--downloads", default=str(Path.home() / "Downloads"))
    args = parser.parse_args()
    downloads, output = Path(args.downloads), Path(args.output)
    if output.exists():
        shutil.rmtree(output)
    output.mkdir(parents=True)
    rejects: list[dict[str, object]] = []
    architect_files = prepare_architects(downloads / "All_State_with_bday_Cleaned.xlsx", output, rejects)
    birthday_file = prepare_birthdays(downloads / "Unmatched_Birthday_Contacts.xlsx", output, rejects)
    campaign_files = prepare_campaigns(downloads / "Campaign_Data.xlsx", downloads / "Email_Campaign_data.xlsx", output, args.owner_id, rejects)
    lead_files = prepare_leads(downloads / "Cleaned_Interakt_Leads.xlsx", output, args.owner_id, rejects)
    write_csv(output / "rejected_rows.csv", ["source", "line", "reason", "name"], rejects)
    manifest = {"architects": architect_files, "birthday_contacts": [birthday_file], "campaigns": campaign_files, "leads": lead_files, "rejected_rows": "rejected_rows.csv", "counts": Counter(row["source"] for row in rejects)}
    (output / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
